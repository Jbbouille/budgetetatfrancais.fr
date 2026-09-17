#!/usr/bin/env python3
"""Reconstruit data/cofog.json a partir des fichiers publies par l'Insee.

Usage :
    pip install openpyxl requests
    python scripts/build_data.py

Le script telecharge les tableaux 3.301 a 3.306 des comptes nationaux
(depenses des administrations publiques ventilees par fonction, nomenclature
COFOG) et les convertit en un seul JSON consomme par le site.

Aucune valeur n'est saisie a la main : tout ce qui s'affiche sur le site vient
de ces fichiers. Les chiffres macroeconomiques 2025 (total, deficit, dette)
vivent dans data/meta.json, qui est maintenu a la main et cite sa source ligne
par ligne.
"""

import datetime
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
    import requests
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("Dependances manquantes : pip install openpyxl requests")

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / ".cache"
OUT = ROOT / "data" / "cofog.json"

# Page Insee : "Depenses des administrations publiques ventilees par fonction
# en 2024" -- https://www.insee.fr/fr/statistiques/8574707
BASE = "https://www.insee.fr/fr/statistiques/fichier/8574707"

TABLES = {
    "3301": "apu",     # S13   -- ensemble des administrations publiques
    "3303": "etat",    # S13111 -- Etat
    "3305": "local",   # S1313  -- administrations publiques locales
    "3306": "secu",    # S1314  -- administrations de securite sociale
}

# Codes d'operation comptable a extraire en plus du total (colonne A du fichier).
NATURES = {
    "D1": "Remuneration des salaries",
    "P2": "Consommations intermediaires",
    "D3": "Subventions",
    "D4": "Interets et revenus de la propriete",
    "D62": "Prestations sociales en especes",
    "D632": "Prestations sociales en nature",
    "D7": "Autres transferts courants",
    "D9": "Transferts en capital",
    "P5L": "Investissement (actifs non financiers)",
    "OED": "Impots sur la production et le revenu",
}


def download(table: str) -> Path:
    CACHE.mkdir(exist_ok=True)
    path = CACHE / f"T_{table}.xlsx"
    if not path.exists():
        url = f"{BASE}/T_{table}.xlsx"
        print(f"  telechargement {url}")
        r = requests.get(url, timeout=120)
        r.raise_for_status()
        path.write_bytes(r.content)
    return path


def read(path: Path, sheet: str, want: str = "OTE"):
    """Lit un bloc d'operation comptable du tableau.

    Le fichier Insee empile plusieurs blocs dans la meme feuille : le total
    (OTE) puis chaque nature de depense. Le code d'operation n'est porte que
    par la premiere ligne du bloc, d'ou la propagation de `cur`.
    """
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    years = [int(c) for c in rows[3][4:] if c is not None and str(c).strip().isdigit()]
    out, cur = {}, None
    # enumerate a partir de 6 : rows[5] est la 6e ligne du classeur, et Excel
    # numerote ses lignes a partir de 1. On retient ce numero pour pouvoir
    # renvoyer l'utilisateur vers la cellule exacte du fichier.
    for excel_row, row in enumerate(rows[5:], start=6):
        head = row[0]
        if isinstance(head, str) and head.strip() and len(head.strip()) <= 5:
            cur = head.strip()
        if cur != want:
            continue
        code = row[2]
        if not isinstance(code, str) or not (code.startswith("GF") or code == "_Z"):
            continue
        label = re.sub(r"^\s*\d{2}(\.\d)?\s*-\s*", "", (row[3] or "")).strip()
        values = [
            round(float(c), 1) if isinstance(c, (int, float)) else None
            for c in row[4 : 4 + len(years)]
        ]
        if values:
            out.setdefault(code, {"label": label, "v": values, "r": excel_row})
    return years, out


def main() -> None:
    print("Construction de data/cofog.json")
    paths = {t: download(t) for t in TABLES}

    years, apu = read(paths["3301"], "T_3301")

    natures = {}
    for sto, label in NATURES.items():
        _, block = read(paths["3301"], "T_3301", want=sto)
        natures[sto] = {
            "label": label,
            "v": {k: v["v"][-1] for k, v in block.items() if len(k) == 4 or k == "_Z"},
        }

    secteurs = {}
    for table, key in TABLES.items():
        if key == "apu":
            continue
        _, block = read(paths[table], f"T_{table}")
        secteurs[key] = {
            k: v["v"] for k, v in block.items() if len(k) == 4 or k == "_Z"
        }

    # Les annees commencent en colonne E (5e colonne) : on precalcule la
    # lettre de colonne de chaque annee pour l'afficher telle quelle.
    colonnes = [get_column_letter(5 + i) for i in range(len(years))]

    data = {
        "genere": datetime.date.today().isoformat(),
        "millesime": years[-1],
        "years": years,
        "source": {
            "fichier": "T_3301.xlsx",
            "feuille": "T_3301",
            "url": f"{BASE}/T_3301.xlsx",
            "page": "https://www.insee.fr/fr/statistiques/8574707",
            "colonnes": colonnes,
        },
        "apu": apu,
        "natures": natures,
        "secteurs": secteurs,
    }

    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(
        json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )

    total = apu["_Z"]["v"][-1]
    somme = sum(apu[k]["v"][-1] for k in apu if len(k) == 4)
    print(f"  annees {years[0]}-{years[-1]}")
    print(f"  total {years[-1]} : {total} Md EUR")
    print(f"  controle : somme des 10 fonctions = {somme:.1f} Md EUR")
    if abs(total - somme) > 0.15:
        sys.exit("ECHEC : la somme des fonctions ne retombe pas sur le total")
    print(f"  ecrit {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
